"""The overseer report and the zonal monthly report: figures, rules and the API."""
import datetime
import io
from decimal import Decimal
from unittest import mock

from django.test import TestCase
from openpyxl import load_workbook
from pypdf import PdfReader
from rest_framework.test import APITestCase

from accounts.models import Role, RolePermission, User
from attendance.models import AttendanceSession, Fellowship, MeetingType
from core.models import Location
from finance.models import Expense, ExpenseCategory, Fund, Giving, PaymentMethod, Remittance, RemittanceLine
from goals.models import Goal
from members.models import Household, Member
from newcomers.models import Newcomer, NewcomerSource

from .church_data import gather, month_list, period_words, whole
from .models import Report, Testimony, Service
from .overseer import build
from .zonal import build_xlsx, plan

D = datetime.date


def text(pdf_bytes):
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf_bytes)).pages)


class Church(TestCase):
    """Three months of a small church: Fridays, a Wednesday that gave way
    to GCK, a workers' meeting, a house fellowship, giving and expenses."""

    @classmethod
    def setUpTestData(cls):
        cls.bh = Location.objects.create(id="bahrain", name="Bahrain HQ", is_core=True)
        cls.qa = Location.objects.create(id="others", name="Qatar")
        cls.fri = MeetingType.objects.create(id="fri-worship", name="Friday Worship Service", day="Friday",
                                             frequency="weekly", detail_level="detailed")
        cls.wed = MeetingType.objects.create(id="wed-rev", name="Wednesday Revival and Evangelism Training", day="Wednesday",
                                             frequency="weekly", detail_level="detailed")
        cls.sat = MeetingType.objects.create(id="sat-workers", name="Saturday Workers Meeting", day="Saturday",
                                             frequency="weekly", detail_level="simple", audience="workers")
        cls.gck = MeetingType.objects.create(id="gck", name="Global Crusade with Kumuyi (GCK)", frequency="occasional",
                                             detail_level="detailed")
        cls.hcf = MeetingType.objects.create(id="fri-house", name="Friday House Caring Fellowship", day="Friday",
                                             frequency="weekly", detail_level="detailed")
        cls.fel = Fellowship.objects.create(name="Sisters", meeting_type=cls.hcf, location=cls.bh)

        def sess(mt, d, loc=None, **kw):
            return AttendanceSession.objects.create(meeting_type=mt, date=d, location=loc or cls.bh, mode="in-person",
                                                    status="filled", **kw)
        # July: two Fridays (26 and 15 people), August: one Friday of 31.
        cls.f1 = sess(cls.fri, D(2026, 7, 17), men=14, women=6, youth_boys=1, youth_girls=1, children_boys=1, children_girls=3, new_comers=4)
        cls.f2 = sess(cls.fri, D(2026, 7, 24), men=6, women=3, youth_boys=1, youth_girls=1, children_boys=1, children_girls=3, new_comers=1)
        cls.f3 = sess(cls.fri, D(2026, 8, 7), men=11, women=7, online_men=2, youth_boys=1, youth_girls=1, children_boys=2, children_girls=7)
        sess(cls.wed, D(2026, 7, 22), men=5, women=2, youth_boys=1, youth_girls=1, children_boys=2, children_girls=5, new_comers=1)
        # Wednesday 29 July gave way to GCK, held 29 to 31 July with a theme.
        for i, d in enumerate((D(2026, 7, 29), D(2026, 7, 30), D(2026, 7, 31))):
            sess(cls.gck, d, men=8 + i, women=4, edition_name="The Season of Divine Turnaround")
        sess(cls.gck, D(2026, 9, 20), men=5, women=4, edition_place="Pakistan")
        sess(cls.sat, D(2026, 7, 18), men=5, women=3)
        cls.h1 = sess(cls.hcf, D(2026, 7, 24), fellowship=cls.fel, women=6)
        # A Qatar Friday, which a Bahrain HQ only report leaves out.
        sess(cls.fri, D(2026, 7, 17), loc=cls.qa, men=9, women=9)
        # Not filled, so not counted.
        AttendanceSession.objects.create(meeting_type=cls.fri, date=D(2026, 7, 31), location=cls.bh, mode="in-person", men=99)

        cash = PaymentMethod.objects.create(name="Cash")
        cls.tithe, cls.offering, cls.dlicc = (Fund.objects.create(name=n) for n in ("Tithe", "Offering", "DLICC vow"))
        Giving.objects.create(date=D(2026, 7, 17), fund=cls.tithe, method=cash, amount=Decimal("950.000"), location=cls.bh)
        Giving.objects.create(date=D(2026, 7, 17), fund=cls.offering, method=cash, amount=Decimal("32.070"), location=cls.bh, session=cls.f1)
        Giving.objects.create(date=D(2026, 7, 24), fund=cls.offering, method=cash, amount=Decimal("9.000"), location=cls.bh, session=cls.f2)
        Giving.objects.create(date=D(2026, 7, 24), fund=cls.offering, method=cash, amount=Decimal("5.500"), location=cls.bh, session=cls.h1)
        Giving.objects.create(date=D(2026, 7, 30), fund=cls.dlicc, method=cash, amount=Decimal("20.000"), location=cls.bh)
        Giving.objects.create(date=D(2026, 8, 7), fund=cls.offering, method=cash, amount=Decimal("49.140"), location=cls.bh, session=cls.f3)
        Giving.objects.create(date=D(2026, 7, 17), fund=cls.offering, method=cash, amount=Decimal("100.000"), location=cls.qa)
        rent, admin = ExpenseCategory.objects.create(name="Rent"), ExpenseCategory.objects.create(name="Church administration")
        Expense.objects.create(date=D(2026, 7, 3), category=rent, amount=Decimal("100.000"), location=cls.bh)
        Expense.objects.create(date=D(2026, 7, 9), category=admin, amount=Decimal("774.910"), location=cls.bh)
        Expense.objects.create(date=D(2026, 8, 3), category=rent, amount=Decimal("200.000"), location=cls.bh)
        r = Remittance.objects.create(month=D(2026, 7, 1), location=cls.bh)
        RemittanceLine.objects.create(remittance=r, fund=cls.dlicc, amount_due=Decimal("20"), amount_sent=Decimal("20"), destination="lagos")

        hh = Household.objects.create(name="Afari")
        for i in range(3):
            Member.objects.create(surname=f"M{i}", first_name="A", location=cls.bh, joined_date=D(2026, 7, 17), household=hh if i < 2 else None,
                                  category=["Worker", "Worker in Training", "General Member"][i])
        Member.objects.create(surname="Q", first_name="A", location=cls.qa, joined_date=D(2026, 7, 17))
        src = NewcomerSource.objects.create(name="Friend")
        for st in ("new", "attending", "member", "not-interested"):
            Newcomer.objects.create(name=f"N {st}", source=src, stage=st, location=cls.bh, created_at=D(2026, 7, 17), stage_since=D(2026, 7, 17))
        Testimony.objects.create(date=D(2026, 8, 9), service=Service.objects.create(name="Friday"), text="God healed my wife.")
        Goal.objects.create(horizon="Short-term", name="At least 10 new visitors a month", target=10, tracking="manual",
                            source="x", target_date="From October 2026", measured_by="Newcomers registered each month")
        Goal.objects.create(horizon="Long-term", name="A church in a new location", target=1, tracking="manual", source="x")


class Figures(Church):
    def test_period_words_and_titles(self):
        self.assertEqual(period_words(month_list((2026, 9), (2026, 9))), ("Monthly", "month", "September 2026"))
        self.assertEqual(period_words(month_list((2026, 7), (2026, 9))), ("Quarterly", "quarter", "July to September 2026"))
        self.assertEqual(period_words(month_list((2026, 11), (2027, 1)))[2], "November 2026 to January 2027")
        self.assertEqual(period_words(month_list((2026, 1), (2026, 12)))[:2], ("Annual", "year"))
        self.assertEqual(period_words(month_list((2026, 7), (2026, 10)))[:2], ("", "period"))

    def test_whole_numbers_round_half_up(self):
        self.assertEqual((whole(20.5), whole(20.49), whole(-2.5)), (21, 20, -3))

    def test_friday_figures_count_filled_sessions_only_with_online_people(self):
        ctx = gather(month_list((2026, 7), (2026, 8)), self.bh)
        self.assertEqual([s.total for s in ctx["fridays"]], [26, 15, 31])
        self.assertEqual(ctx["friday_average"], 24)          # 72 / 3
        self.assertEqual(ctx["newcomer_visits"], 6)

    def test_whole_church_includes_every_location(self):
        ctx = gather(month_list((2026, 7), (2026, 7)), None)
        self.assertEqual(len(ctx["fridays"]), 3)
        self.assertEqual(ctx["income"], Decimal("1116.570"))

    def test_money_by_month_fund_and_category(self):
        ctx = gather(month_list((2026, 7), (2026, 8)), self.bh)
        jul, aug = ctx["months"]
        self.assertEqual(jul["income"], Decimal("1016.570"))
        self.assertEqual(jul["by_fund"]["Offering"], Decimal("46.570"))
        self.assertEqual(jul["expenses"], Decimal("874.910"))
        self.assertEqual(aug["net"], Decimal("-150.860"))
        self.assertEqual(ctx["categories"], ["Church administration", "Rent"])

    def test_meeting_averages_and_programmes(self):
        ctx = gather(month_list((2026, 7), (2026, 9)), self.bh)
        self.assertEqual(list(ctx["meetings"])[0], "Friday Worship Service")
        self.assertEqual(ctx["meetings"]["Friday Worship Service"][(2026, 7)], (21, 2))
        progs = ctx["programmes"]
        self.assertEqual([(p["label"], p["dates"], p["days"], p["average"]) for p in progs], [
            ('GCK, "The Season of Divine Turnaround"', "29 to 31 July", 3, 13),
            ("GCK, Pakistan", "20 September", 1, 9)])

    def test_people_as_they_stand(self):
        p = gather(month_list((2026, 7), (2026, 7)), self.bh)["people"]
        self.assertEqual((p["members"], p["newcomers"], p["congregation"], p["households"], p["workers"], p["in_training"]),
                         (3, 2, 5, 1, 1, 1))
        self.assertIsNone(p["contacts"])           # contacts have no location
        self.assertEqual(gather(month_list((2026, 7), (2026, 7)), None)["people"]["contacts"], 0)


class Zonal(Church):
    def tables(self, months, loc=None):
        p = plan(gather(month_list(*months), loc))
        return p, {t["title"]: t for s in p["sections"] for t in s["tables"] if "title" in t}

    def test_gck_marks_the_wednesday_it_replaced(self):
        p, t = self.tables(((2026, 7), (2026, 7)), self.bh)
        rows = t["Weekly fellowship attendance, July 2026"]["rows"]
        self.assertIn(["29 Jul", "GCK", "", "", "", "", "", "", ""], rows)
        self.assertTrue(p["any_marker"])
        # Friday Worship first, house fellowships left out.
        self.assertEqual(rows[0], ["Friday Worship Service"])
        self.assertNotIn(["Friday House Caring Fellowship"], rows)
        # The Friday on 31 July was left unrecorded while GCK was on.
        self.assertIn(["31 Jul", "GCK", "", "", "", "", "", "", ""], rows)

    def test_giving_by_service_adds_up_to_the_month(self):
        _, t = self.tables(((2026, 7), (2026, 7)), self.bh)
        g = t["Giving by service, July 2026 (BHD)"]
        self.assertEqual(g["header"], ["Date", "DLICC vow", "Tithe", "Offering", "Total"])
        self.assertIn(["17 Jul", "0.00", "950.00", "32.07", "982.07"], g["rows"])   # tithe without a session goes to that Friday
        self.assertEqual(g["rows"][-1], ["Total", "20.00", "950.00", "46.57", "1,016.57"])
        others = [r[0] for r in g["rows"] if "," in r[0]]
        self.assertEqual(sorted(others), ["24 Jul, Friday House Caring Fellowship (Sisters)", "30 Jul, Not linked to a meeting"])

    def test_workers_special_and_financial_summary(self):
        _, t = self.tables(((2026, 7), (2026, 7)), self.bh)
        self.assertEqual(t["Workers' and leaders' meetings, July 2026"]["rows"], [["Saturday Workers Meeting"], ["18 Jul", "5", "3", "8"]])
        g = t["Global Crusade with Kumuyi (GCK), July 2026"]["rows"]
        self.assertEqual(g[0][:3], ["29 Jul", "GCK, The Season of Divine Turnaround", "8"])
        self.assertEqual(g[-3:], [["Average", "", "", "", "", "", "", "", "", "13"], ["Highest", "", "", "", "", "", "", "", "", "14"],
                                  ["Lowest", "", "", "", "", "", "", "", "", "12"]])
        f = {r[0]: r for r in t["Financial summary, July 2026 (BHD)"]["rows"]}
        self.assertEqual(f["DLICC project"][1:], ["Sent to the DLBC account in Lagos", "20.00"])
        self.assertEqual(f["House caring fellowship leaders' offerings"][1:], ["From 1 house caring fellowship meeting", "5.50"])
        self.assertEqual(f["Total cash balance"][2], "121.66")       # 1016.57 - 874.91 - 20
        _, t8 = self.tables(((2026, 8), (2026, 8)), self.bh)
        f8 = {r[0]: r for r in t8["Financial summary, August 2026 (BHD)"]["rows"]}
        self.assertEqual(f8["Offering HQ"][1], "No remittance recorded for this month")

    def test_summary_only_for_several_months_and_spreadsheet_matches(self):
        p1, _ = self.tables(((2026, 7), (2026, 7)), self.bh)
        self.assertIsNone(p1["summary"])
        ctx = gather(month_list((2026, 7), (2026, 8)), self.bh)
        p = plan(ctx)
        self.assertEqual(p["summary"]["tables"][0]["rows"][-1], ["Total", "4", "88", "1,065.71", "1,074.91", "-29.20"])
        wb = load_workbook(io.BytesIO(build_xlsx(ctx)))
        self.assertEqual(wb.sheetnames, ["July 2026", "August 2026", "Summary"])
        values = [c.value for row in wb["Summary"].iter_rows() for c in row]
        self.assertIn("Table 13. Monthly totals, July to August 2026.", values)
        self.assertIn(1065.71, values)


class OverseerPdf(Church):
    def test_one_month_is_the_monthly_report_without_month_tables(self):
        pdf, title = build(gather(month_list((2026, 8), (2026, 8)), self.bh), {})
        self.assertEqual(title, "DCLM Bahrain Monthly Report, August 2026")
        t = text(pdf)
        self.assertIn("The Month at a Glance", t)
        self.assertNotIn("Average attendance by meeting and month", t)
        self.assertNotIn("Summary", t.split("Appendix")[0].replace("Monthly Report", ""))   # empty box, no section
        self.assertIn("Net shortfall (BHD)", t)

    def test_quarter_with_written_parts_and_goals(self):
        texts = {"summary": "By the grace of God, a good quarter.", "challenges": "Members live far apart.\n- Transport plan",
                 "training": "1. Training lasts 2 months.", "conclusion": "We end well."}
        pdf, title = build(gather(month_list((2026, 7), (2026, 9)), self.bh), texts)
        self.assertEqual(title, "DCLM Bahrain Quarterly Report, July to September 2026")
        cover = " ".join(PdfReader(io.BytesIO(pdf)).pages[0].extract_text().split())
        self.assertIn("Quarterly Report July to September 2026 THE QUARTER AT A GLANCE", cover)
        long, _ = build(gather(month_list((2026, 1), (2026, 10)), None), {})
        cover = " ".join(PdfReader(io.BytesIO(long)).pages[0].extract_text().split())
        self.assertIn("Report January to October 2026 THE PERIOD AT A GLANCE", cover)
        lt = " ".join(text(long).split())
        self.assertIn("Giving by fund and month (BHD), January to June 2026.", lt)
        self.assertIn("Giving by fund and month (BHD), July to October 2026.", lt)
        t = " ".join(text(pdf).split())
        for s in ("1. The Quarter at a Glance", "2. Summary", "By the grace of God, a good quarter.", "Average attendance by meeting and month",
                  "Special programmes", "Giving by fund and month (BHD)",
                  "Key ratios by month", "Membership and Outreach", "God healed my wife.", "Challenges and Responses", "Transport plan",
                  "Short-term goals", "From October 2026", "Newcomers registered each month", "Long-term goals", "Not set",
                  "Training lasts 2 months.", "We end well.", "Appendix: Each Month in Detail", "Workers 1, with 1 in training"):
            self.assertIn(s, t)
        self.assertNotIn("Contacts yet to visit", t)          # one location's report
        self.assertIn("Friday Worship attendances 41 (2 services) 31 (1 service) None", t)
        pdf, _ = build(gather(month_list((2026, 7), (2026, 8)), None), {})
        t = " ".join(text(pdf).split())
        self.assertIn("Change in average attendance from July to August", t)
        self.assertIn("From July to August, Friday Worship Service grew.", t)
        self.assertIn("Contacts yet to visit 0", t)


class Api(Church):
    def setUp(self):
        from rest_framework.test import APIClient
        self.client = APIClient()
        self.role = Role.objects.create(name="Administrator")
        RolePermission.objects.create(role=self.role, module="reports", can_view=True, can_create=True, can_edit=True, can_delete=True)
        self.vrole = Role.objects.create(name="Viewer")
        RolePermission.objects.create(role=self.vrole, module="reports", can_view=True)
        self.admin = User.objects.create_user(email="a@t.com", password="x", role=self.role)
        self.viewer = User.objects.create_user(email="v@t.com", password="x", role=self.vrole)
        self.qatar = User.objects.create_user(email="q@t.com", password="x", role=self.role, location=self.qa)
        self.nobody = User.objects.create_user(email="n@t.com", password="x", role=Role.objects.create(name="None"))
        self.today = mock.patch("django.utils.timezone.localdate", return_value=D(2026, 10, 13))
        self.today.start()

    def tearDown(self):
        self.today.stop()

    def post(self, user, url, **data):
        self.client.force_authenticate(user)
        return self.client.post(url, data, format="json")

    def test_overseer_for_three_months_then_brought_forward(self):
        r = self.post(self.admin, "/api/reports/generate/", kind="overseer", from_year=2026, from_month=7, to_year=2026, to_month=9,
                      location="bahrain", summary="  Summary text.  ", challenges="Hard things.")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual((r.data["kind_label"], r.data["period_label"], r.data["summary"]), ("Overseer report", "July to September 2026", "Summary text."))
        rep = Report.objects.get()
        self.assertIn("DCLM-Bahrain-Quarterly-Report-July-to-September-2026-bahrain", rep.pdf_file.name)
        again = self.post(self.admin, "/api/reports/generate/", kind="overseer", from_year=2026, from_month=7, to_year=2026, to_month=9, location="bahrain")
        self.assertEqual(again.status_code, 409)
        self.client.force_authenticate(self.admin)
        last = self.client.get("/api/reports/last-text/?location=bahrain").data
        self.assertEqual((last["summary"], last["challenges"], last["from_report"]), ("Summary text.", "Hard things.", "July to September 2026"))
        self.assertEqual(self.client.get("/api/reports/last-text/").data["summary"], "")   # the whole church has none yet

    def test_one_month_and_zonal_are_different_reports(self):
        a = self.post(self.admin, "/api/reports/generate/", kind="overseer", from_year=2026, from_month=9, to_year=2026, to_month=9)
        b = self.post(self.admin, "/api/reports/generate/", kind="zonal", from_year=2026, from_month=9, to_year=2026, to_month=9)
        c = self.post(self.admin, "/api/reports/generate/", period_year=2026, period_month=8)     # older screens
        self.assertEqual([a.status_code, b.status_code, c.status_code], [201, 201, 201])
        self.assertEqual(Report.objects.get(pk=c.data["id"]).first_month, (2026, 8))

    def test_months_are_checked(self):
        for body, msg in (
            (dict(from_year=2026, from_month=9, to_year=2026, to_month=7), "before the last"),
            (dict(from_year=2026, from_month=10, to_year=2026, to_month=11), "not started"),
            (dict(from_year=2025, from_month=1, to_year=2026, to_month=1), "up to 12"),
        ):
            r = self.post(self.admin, "/api/reports/generate/", kind="overseer", **body)
            self.assertEqual(r.status_code, 400)
            self.assertIn(msg, str(r.data))

    def test_preview_saves_nothing_and_viewers_may_preview(self):
        r = self.post(self.viewer, "/api/reports/preview/", kind="zonal", from_year=2026, from_month=7, to_year=2026, to_month=8)
        self.assertEqual((r.status_code, r["Content-Type"]), (200, "application/pdf"))
        z = text(r.content)
        self.assertIn("Monthly Reports, July to August 2026", z)
        self.assertNotIn("About th", z)
        self.assertNotIn("Amounts are in BHD", z)
        self.assertEqual(Report.objects.count(), 0)
        self.assertEqual(self.post(self.viewer, "/api/reports/generate/", kind="zonal", from_year=2026, from_month=7, to_year=2026, to_month=7).status_code, 403)
        self.assertEqual(self.post(self.nobody, "/api/reports/preview/", kind="zonal", from_year=2026, from_month=7, to_year=2026, to_month=7).status_code, 403)

    def test_location_limited_person_gets_their_own_figures(self):
        # Asking for another location's figures is refused outright.
        r = self.post(self.qatar, "/api/reports/generate/", kind="overseer", from_year=2026, from_month=7, to_year=2026, to_month=7, location="bahrain")
        self.assertEqual(r.status_code, 403)
        r = self.post(self.qatar, "/api/reports/generate/", kind="overseer", from_year=2026, from_month=7, to_year=2026, to_month=7)
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data["location_name"], "Qatar")
        rep = Report.objects.get()
        rep.pdf_file.open("rb")
        t = text(rep.pdf_file.read())
        rep.pdf_file.close()
        self.assertIn("The figures cover Qatar only.", " ".join(t.split()))
        self.assertIn("18 (1 service)", t)

    def test_zonal_spreadsheet_for_a_range(self):
        self.client.force_authenticate(self.viewer)
        r = self.client.get("/api/reports/spreadsheet/?from=2026-07&to=2026-09&location=bahrain")
        self.assertEqual(r.status_code, 200)
        self.assertIn('DCLM-Bahrain-Zonal-2026-07-to-2026-09-Bahrain-HQ.xlsx', r["Content-Disposition"])
        self.assertEqual(load_workbook(io.BytesIO(r.content)).sheetnames, ["July 2026", "August 2026", "September 2026", "Summary"])
        old = self.client.get("/api/reports/spreadsheet/?year=2026&month=8")
        self.assertEqual(load_workbook(io.BytesIO(old.content)).sheetnames, ["August 2026"])
        self.assertEqual(self.client.get("/api/reports/spreadsheet/?from=2026-12").status_code, 400)


class Goals(APITestCase):
    def test_target_date_and_measured_by_can_be_edited(self):
        role = Role.objects.create(name="Administrator")
        RolePermission.objects.create(role=role, module="goals", can_view=True, can_create=True, can_edit=True, can_delete=True)
        user = User.objects.create_user(email="g@t.com", password="x", role=role)
        self.client.force_authenticate(user)
        r = self.client.post("/api/goals/", {"horizon": "Short-term", "name": "Raise Friday average", "target": 30, "tracking": "manual",
                                             "source": "Manual", "target_date": "December 2026", "measured_by": "Friday register"}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        r = self.client.patch(f"/api/goals/{r.data['id']}/", {"name": "Raise Friday average to 30", "measured_by": "Attendance register"}, format="json")
        self.assertEqual((r.data["name"], r.data["target_date"], r.data["measured_by"]), ("Raise Friday average to 30", "December 2026", "Attendance register"))
        from accounts.models import AuditLog
        self.assertTrue(AuditLog.objects.filter(action="Edited", entity_type="Goal").exists())


class NothingRecorded(TestCase):
    def test_both_reports_still_come_out(self):
        loc = Location.objects.create(id="bahrain", name="Bahrain HQ", is_core=True)
        for location in (None, loc):
            ctx = gather(month_list((2026, 1), (2026, 2)), location)
            pdf, _ = build(ctx, {})
            t = " ".join(text(pdf).split())
            self.assertIn("No Friday Worship Service was recorded", t)
            self.assertIn("No giving recorded in the period", t)
            self.assertIn("No testimonies were recorded", t)
            from .zonal import build_pdf
            z = " ".join(text(build_pdf(ctx)[0]).split())
            self.assertIn("No weekly services were recorded in January 2026.", z)
            cover = " ".join(PdfReader(io.BytesIO(build_pdf(ctx)[0])).pages[0].extract_text().split())
            self.assertEqual("Bahrain HQ only" in cover, location is not None)
            self.assertIn("Total cash balance To be retained in", z)
            load_workbook(io.BytesIO(build_xlsx(ctx)))

"""
The zonal monthly report: the layout Zonal asks for, one section per
month, filled in entirely from the system. The same tables feed the PDF
and the spreadsheet, so the two can never disagree.
"""
import datetime
import io
from collections import OrderedDict, defaultdict

from reportlab.platypus import PageBreak

from . import house as H
from .church_data import (ZERO, FRIDAY, day, female, gather, groups, is_special, is_weekly_service,
                          is_workers, male, short_meeting, whole)

DESTINATIONS = [("dubai", "Offering HQ", "Sent to the DLBC account in Dubai"),
                ("lagos", "DLICC project", "Sent to the DLBC account in Lagos"),
                ("qatar", "GCC special project or GCK", "Sent to the DLBC account in Qatar")]
WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def _order(mt):
    """Friday Worship first, then the week from Monday."""
    return (0 if mt.id == FRIDAY else 1, WEEKDAYS.index(mt.day) if mt.day in WEEKDAYS else 9, mt.name)


def _markers(mf, meeting_types):
    """Dates a weekly meeting did not hold because an occasional meeting
    such as GCK took its place: {meeting id: {date: 'GCK'}}."""
    specials = defaultdict(list)
    for s in mf["sessions"]:
        if is_special(s.meeting_type):
            specials[s.date].append(short_meeting(s.meeting_type.name))
    held = defaultdict(set)
    for s in mf["sessions"]:
        held[s.meeting_type_id].add(s.date)
    out = defaultdict(dict)
    for mt in meeting_types:
        # House fellowships meet in many places, so a missing one is not a
        # service that gave way.
        if mt.day not in WEEKDAYS or mt.fellowship_set.exists():
            continue
        wd = WEEKDAYS.index(mt.day)
        d = mf["start"]
        while d <= mf["end"]:
            if d.weekday() == wd and d not in held[mt.id] and specials.get(d):
                out[mt.id][d] = specials[d][0]
            d += datetime.timedelta(days=1)
    return out


def _grouped(mf, kind, markers):
    """[(meeting type, [(date, session or marker text)])] for weekly
    services or workers' meetings, fellowship meetings left out."""
    test = is_weekly_service if kind == "services" else is_workers
    per = OrderedDict()
    for s in mf["sessions"]:
        if test(s.meeting_type) and not s.fellowship_id:
            per.setdefault(s.meeting_type, []).append((s.date, s))
    types = {s.meeting_type.id: s.meeting_type for s in mf["sessions"]}
    for mid, dates in markers.items():
        mt = types.get(mid) or next((m for m in per if m.id == mid), None)
        if mt is None:
            from attendance.models import MeetingType
            mt = MeetingType.objects.filter(pk=mid).first()
        if mt and test(mt):
            per.setdefault(mt, []).extend((d, text) for d, text in dates.items())
    return [(mt, sorted(rows, key=lambda r: r[0])) for mt, rows in sorted(per.items(), key=lambda x: _order(x[0]))]


def _remitted(scope, y, m):
    from finance.models import Remittance
    qs = scope(Remittance.objects.filter(month__year=y, month__month=m)).prefetch_related("lines")
    sent, any_ = defaultdict(lambda: ZERO), False
    for r in qs:
        any_ = True
        for line in r.lines.all():
            sent[line.destination] += line.amount_sent
    return sent, any_


def plan(ctx):
    """Sections of tables, shared by the PDF and the spreadsheet. Each
    table: dict(title, header, rows, widths, bands, totals)."""
    from attendance.models import MeetingType
    months = ctx["months"]
    scope = ctx["scope"]
    weekly = [mt for mt in MeetingType.objects.filter(frequency="weekly")]
    used = {s.meeting_type_id for s in ctx["sessions"]}
    weekly = [mt for mt in weekly if mt.id in used]
    funds = sorted(ctx["funds"], key=lambda f: f.lower() in ("tithe", "offering"))  # projects first, as Zonal lists them
    cats = ctx["categories"]
    where = "the location" if scope.location else "the church"
    sections, summary_rows, any_marker = [], [], False
    for mf in months:
        tables = []
        markers = _markers(mf, weekly)
        any_marker = any_marker or any(markers.values())
        name = mf["name"]

        # weekly services
        services = _grouped(mf, "services", markers)
        rows, bands = [], set()
        stats = []
        weekly_count = weekly_people = 0
        for mt, items in services:
            bands.add(len(rows))
            rows.append([mt.name])
            totals = []
            for d, s in items:
                if isinstance(s, str):
                    rows.append([day(d), s, "", "", "", "", "", "", ""])
                    continue
                g = groups(s)
                rows.append([day(d), str(g["men"]), str(g["women"]), str(g["youth_boys"]), str(g["youth_girls"]),
                             str(g["children_boys"]), str(g["children_girls"]), str(s.new_comers), str(s.total)])
                totals.append(s.total)
            if totals:
                stats.append([mt.name, str(len(totals)), str(whole(sum(totals) / len(totals))), str(max(totals)), str(min(totals))])
                weekly_count += len(totals)
                weekly_people += sum(totals)
        header = ["Date", "Men", "Women", "Youth boys", "Youth girls", "Children boys", "Children girls", "Newcomers", "Total"]
        w = [0.09] + [0.11] * 6 + [0.13, 0.12]
        if rows:
            tables.append(dict(title=f"Weekly fellowship attendance, {name}", header=header, rows=rows, widths=w, bands=bands, totals=set()))
            tables.append(dict(title=f"Attendance summary, {name}", header=["Meeting", "Services", "Average", "Highest", "Lowest"],
                               rows=stats, widths=[0.48, 0.13, 0.13, 0.13, 0.13], bands=set(), totals=set()))
        else:
            tables.append(dict(note=f"No weekly services were recorded in {name}."))

        # giving by service
        service_sessions = [s for _, items in services for _, s in items if not isinstance(s, str)]
        by_session = defaultdict(lambda: defaultdict(lambda: ZERO))
        others = OrderedDict()
        friday_first = sorted(service_sessions, key=lambda s: (s.date, s.meeting_type_id != FRIDAY))
        for gv in mf["giving"]:
            target = None
            if gv.session_id and any(s.id == gv.session_id for s in service_sessions):
                target = gv.session_id
            elif not gv.session_id:
                same = [s for s in friday_first if s.date == gv.date and s.location_id == gv.location_id]
                target = same[0].id if same else None
            if target:
                by_session[target][gv.fund.name] += gv.amount
            else:
                label = (gv.session.meeting_type.name if gv.session_id else "Not linked to a meeting")
                if gv.session_id and gv.session.fellowship_id:
                    label = f"{gv.session.meeting_type.name} ({gv.session.fellowship.name})"
                others.setdefault((gv.date, label), defaultdict(lambda: ZERO))[gv.fund.name] += gv.amount
        cols = funds or []
        rows, bands, col_tot = [], set(), defaultdict(lambda: ZERO)
        for mt, items in services:
            sess = [s for _, s in items if not isinstance(s, str)]
            if not sess:
                continue
            bands.add(len(rows))
            rows.append([mt.name])
            for s in sess:
                vals = [by_session[s.id].get(f, ZERO) for f in cols]
                rows.append([day(s.date)] + [H.money(v) for v in vals] + [H.money(sum(vals, ZERO))])
                for f, v in zip(cols, vals):
                    col_tot[f] += v
        if others:
            bands.add(len(rows))
            rows.append(["Other giving"])
            for (d, label), per in others.items():
                vals = [per.get(f, ZERO) for f in cols]
                rows.append([f"{day(d)}, {label}"] + [H.money(v) for v in vals] + [H.money(sum(vals, ZERO))])
                for f, v in zip(cols, vals):
                    col_tot[f] += v
        if rows:
            rows.append(["Total"] + [H.money(col_tot[f]) for f in cols] + [H.money(mf["income"])])
            first = 0.26 if others else 0.14
            tables.append(dict(title=f"Giving by service, {name} (BHD)", header=["Date"] + cols + ["Total"], rows=rows,
                               widths=[first] + [(1 - first) / (len(cols) + 1)] * (len(cols) + 1), bands=bands, totals={len(rows) - 1}))

        # workers' and leaders' meetings
        workers = _grouped(mf, "workers", markers)
        rows, bands = [], set()
        for mt, items in workers:
            bands.add(len(rows))
            rows.append([mt.name])
            for d, s in items:
                rows.append([day(d), s, "", ""] if isinstance(s, str) else [day(d), str(male(s)), str(female(s)), str(s.total)])
        if rows:
            tables.append(dict(title=f"Workers' and leaders' meetings, {name}", header=["Date", "Male", "Female", "Total"],
                               rows=rows, widths=[0.34, 0.22, 0.22, 0.22], bands=bands, totals=set()))

        # occasional meetings, one table each
        specials = OrderedDict()
        for s in mf["sessions"]:
            if is_special(s.meeting_type):
                specials.setdefault(s.meeting_type.name, []).append(s)
        for mname, sess in sorted(specials.items(), key=lambda x: (not x[0].startswith("Global Crusade"), x[0])):
            rows = []
            for s in sess:
                g = groups(s)
                prog = short_meeting(mname) + (f", {s.edition_name}" if s.edition_name else (f", {s.edition_place}" if s.edition_place else ""))
                rows.append([day(s.date), prog, str(g["men"]), str(g["women"]), str(g["youth_boys"]), str(g["youth_girls"]),
                             str(g["children_boys"]), str(g["children_girls"]), str(s.new_converts), str(s.total)])
            tt = [s.total for s in sess]
            n0 = len(rows)
            for lab, v in (("Average", whole(sum(tt) / len(tt))), ("Highest", max(tt)), ("Lowest", min(tt))):
                rows.append([lab, "", "", "", "", "", "", "", "", str(v)])
            tables.append(dict(title=f"{mname}, {name}",
                               header=["Date", "Programme", "Men", "Women", "Youth boys", "Youth girls", "Children boys", "Children girls", "New converts", "Total"],
                               rows=rows, widths=[0.08, 0.20] + [0.09] * 8, bands=set(), totals=set(range(n0, n0 + 3))))

        # expenditures
        rows = [[str(i + 1), c, H.money(mf["by_category"].get(c, ZERO))] for i, c in enumerate(cats)]
        if rows:
            rows.append(["", "Total", H.money(mf["expenses"])])
            tables.append(dict(title=f"Expenditures, {name} (BHD)", header=["Number", "Description", "Amount"], rows=rows,
                               widths=[0.14, 0.56, 0.30], bands=set(), totals={len(rows) - 1}))

        # financial summary
        sent, recorded = _remitted(scope, mf["year"], mf["month"])
        hcf = [gv for gv in mf["giving"] if gv.session_id and gv.session.fellowship_id]
        hcf_total = sum((gv.amount for gv in hcf), ZERO)
        hcf_meetings = len({gv.session_id for gv in hcf})
        rows = [["Total income", f"All income of {where}", H.money(mf["income"])],
                ["Total expenses", f"All expenses of {where}", H.money(mf["expenses"])]]
        out_total = ZERO
        for key, label, note in DESTINATIONS:
            v = sent.get(key, ZERO)
            out_total += v
            rows.append([label, note if recorded else "No remittance recorded for this month", H.money(v)])
        rows.append(["House caring fellowship leaders' offerings",
                     f"From {hcf_meetings} house caring fellowship {'meeting' if hcf_meetings == 1 else 'meetings'}" if hcf else "None recorded this month",
                     H.money(hcf_total)])
        balance = mf["income"] - mf["expenses"] - out_total
        rows.append(["Total cash balance", f"To be retained in {where}", H.money(balance)])
        tables.append(dict(title=f"Financial summary, {name} (BHD)", header=["Item", "Note", "Amount"], rows=rows,
                           widths=[0.38, 0.40, 0.22], bands=set(), totals={len(rows) - 1}))
        sections.append({"title": name, "tables": tables})
        summary_rows.append([name, weekly_count, weekly_people, mf["income"], mf["expenses"], balance])

    summary = None
    if len(months) > 1:
        rows = [[r[0], str(r[1]), str(r[2]), H.money(r[3]), H.money(r[4]), H.money(r[5])] for r in summary_rows]
        rows.append(["Total", str(sum(r[1] for r in summary_rows)), str(sum(r[2] for r in summary_rows)),
                     H.money(sum((r[3] for r in summary_rows), ZERO)), H.money(sum((r[4] for r in summary_rows), ZERO)),
                     H.money(sum((r[5] for r in summary_rows), ZERO))])
        noun = ctx["noun"]
        summary = {"title": (f"{noun.title()} Summary" if noun != "period" else "Summary") + f", {ctx['label']}",
                   "tables": [dict(title=f"Monthly totals, {ctx['label']}",
                                   header=["Month", "Weekly services", "Weekly attendance", "Income (BHD)", "Expenses (BHD)", "Cash balance (BHD)"],
                                   rows=rows, widths=[0.2, 0.14, 0.15, 0.17, 0.17, 0.17], bands=set(), totals={len(rows) - 1})]}
    return {"sections": sections, "summary": summary, "any_marker": any_marker}


def build_pdf(ctx):
    p = plan(ctx)
    many = len(ctx["months"]) > 1
    title = "Monthly Reports" if many else "Monthly Report"
    running = f"{title}, {ctx['label']}"
    n = H.Numbers()
    when = f"{ctx['taken_on'].day} {ctx['taken_on']:%B %Y}"

    def cover(c, d):
        # A location's report says so, so it is never read as the whole church's.
        H.draw_cover(c, title, ctx["label"], when, scope=ctx["scope"].label if ctx["scope"].location else "")

    story = H.start_inner()
    for i, sec in enumerate(p["sections"]):
        story += ([PageBreak()] if i else []) + [H.heading(sec["title"])]
        for t in sec["tables"]:
            if "note" in t:
                story.append(H.para(H.safe(t["note"])))
                continue
            story += H.captioned_table(n.table(t["title"]), t["header"], t["rows"], t["widths"], bands=t["bands"], totals=t["totals"])
    if p["summary"]:
        story += [PageBreak(), H.heading(p["summary"]["title"])]
        t = p["summary"]["tables"][0]
        story += H.captioned_table(n.table(t["title"]), t["header"], t["rows"], t["widths"], totals=t["totals"])
        months = ctx["months"]
        names = list(ctx["meetings"].keys())
        shades = [H.colors.HexColor(c) for c in ("#C9DAF2", "#86AEE0", "#2B6FD6", "#1F4E79", "#5B7FA6", "#BEBEB2")]
        series = [(mf["short"], shades[i % len(shades)] if len(months) <= 6 else H.PALETTE[i % 8],
                   [ctx["meetings"][nm].get((mf["year"], mf["month"]), (0, 0))[0] for nm in names]) for i, mf in enumerate(months)]
        if names:
            story.append(H.captioned_figure(n.figure("Average attendance by meeting and month"),
                                            H.grouped_bars([short_label(nm) for nm in names], series, height=200)))
        story.append(H.captioned_figure(n.figure("Giving and expenses by month"), H.grouped_bars(
            [mf["short"] for mf in months],
            [("Giving", H.PALETTE[1], [mf["income"] for mf in months]), ("Expenses", H.PALETTE[3], [mf["expenses"] for mf in months])],
            fmt=H.money, sublabels=[f"Net {H.money(mf['net'])}" for mf in months], height=185)))
    buf = io.BytesIO()
    H.make_doc(buf, f"{H.SHORT} {running}", running, cover).build(story)
    return buf.getvalue(), f"{H.SHORT} {running}"


def short_label(name):
    """Chart labels under the bars: drop the weekday, keep it readable."""
    for wd in WEEKDAYS:
        if name.startswith(wd + " "):
            name = name[len(wd) + 1:]
    name = short_meeting(name) if len(name) > 22 else name
    return name if len(name) <= 22 else name[:21] + "…"


def build_xlsx(ctx):
    """The same tables in a workbook: one sheet per month, and a summary."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    p = plan(ctx)
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="FFBFCBDA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill, zebra_fill, band_fill = (PatternFill("solid", fgColor=c) for c in ("FF1F4E79", "FFE8EFF7", "FFD9D9D9"))
    left = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    n = H.Numbers()

    def sheet(title, tables, intro=None):
        ws = wb.create_sheet(title[:31])
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 2
        r = 2
        c = ws.cell(row=r, column=2, value=f"{H.CHURCH}, {title}")
        c.font = Font(name="Arial", size=13, bold=True, color="FF1F4E79")
        r += 1
        if intro:
            ws.cell(row=r, column=2, value=intro).font = Font(name="Arial", size=9.5, color="FF5A667C")
            r += 1
        widest = defaultdict(int)
        for t in tables:
            if "note" in t:
                r += 1
                ws.cell(row=r, column=2, value=t["note"]).font = Font(name="Arial", size=10)
                r += 1
                continue
            r += 1
            cap = ws.cell(row=r, column=2, value=n.table(t["title"]) + ("" if t["title"].endswith(".") else "."))
            cap.font = Font(name="Arial", size=10, bold=True)
            cap.alignment = Alignment(horizontal="center")
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=1 + len(t["header"]))
            r += 1
            for i, h in enumerate(t["header"]):
                cell = ws.cell(row=r, column=2 + i, value=h)
                cell.font = Font(name="Arial", size=9.5, bold=True, color="FFFFFFFF")
                cell.fill, cell.border, cell.alignment = head_fill, border, left
            ws.row_dimensions[r].height = 30
            r += 1
            zebra = 0
            for k, row in enumerate(t["rows"]):
                band, total = k in t["bands"], k in t["totals"]
                cells = row + [""] * (len(t["header"]) - len(row))
                for i, v in enumerate(cells):
                    val = v
                    if isinstance(v, str) and i > 0 and v.replace(",", "").replace(".", "", 1).lstrip("-").isdigit():
                        val = float(v.replace(",", "")) if "." in v else int(v.replace(",", ""))
                    cell = ws.cell(row=r, column=2 + i, value=val)
                    cell.font = Font(name="Arial", size=9.5, bold=total)
                    cell.border, cell.alignment = border, left
                    if isinstance(val, float):
                        cell.number_format = "#,##0.00"
                    if band or total:
                        cell.fill = band_fill
                    elif zebra % 2 == 1:
                        cell.fill = zebra_fill
                    widest[2 + i] = max(widest[2 + i], len(str(v)) + 3, len(t["header"][i]) // 2 + 4 if i < len(t["header"]) else 0)
                if band:
                    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=1 + len(t["header"]))
                    zebra = 0
                elif not total:
                    zebra += 1
                ws.row_dimensions[r].height = 18
                r += 1
            r += 1
        for col, w in widest.items():
            ws.column_dimensions[get_column_letter(col)].width = min(max(w, 10), 46)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        return ws

    for i, sec in enumerate(p["sections"]):
        sheet(sec["title"], sec["tables"], ctx["scope"].label if (i == 0 and ctx["scope"].location) else None)
    if p["summary"]:
        sheet("Summary", p["summary"]["tables"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render(months, location):
    return build_pdf(gather(months, location))


def spreadsheet(months, location):
    return build_xlsx(gather(months, location))

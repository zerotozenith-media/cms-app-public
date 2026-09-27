"""
Online attendance, fellowships, and the fields the Zonal reporting needs.
"""
import datetime

from django.utils import timezone
from rest_framework.test import APITestCase

from accounts.models import Role, RolePermission, User
from attendance.models import AttendanceSession, Fellowship, MeetingType
from core.models import Location
from finance.models import Fund, Giving, PaymentMethod
from members.models import Member


class Base(APITestCase):
    def setUp(self):
        self.bahrain = Location.objects.create(id="bahrain", name="Bahrain", is_core=True)
        role = Role.objects.create(name="Administrator")
        for module in ["members", "attendance", "newcomers", "finance",
                       "goals", "reports", "outreach", "admin"]:
            RolePermission.objects.create(
                role=role, module=module,
                can_view=True, can_create=True, can_edit=True, can_delete=True)
        self.admin = User.objects.create_user(email="a@t.com", password="x", role=role)
        self.client.force_authenticate(user=self.admin)
        self.mt = MeetingType.objects.create(
            id="fri-worship", name="Friday Worship Service", day="Friday",
            frequency="weekly", detail_level="detailed")
        self.house = MeetingType.objects.create(
            id="fri-house", name="Friday House Caring Fellowship", day="Friday",
            frequency="weekly", detail_level="detailed")
        self.today = timezone.localdate()

    def session(self, meeting=None, **kw):
        kw.setdefault("status", "pending")
        return AttendanceSession.objects.create(
            meeting_type=meeting or self.mt, date=self.today,
            location=self.bahrain, mode="in-person", **kw)


class OnlineAttendanceTestCase(Base):
    """
    Any meeting can be hybrid. A total that counted only the room would
    understate the month, and the church would under-report itself.
    """
    def test_the_total_counts_the_room_and_online_together(self):
        s = self.session(men=10, women=12, online_men=3, online_women=4)
        self.assertEqual(s.in_person_total, 22)
        self.assertEqual(s.online_total, 7)
        self.assertEqual(s.total, 29)

    def test_a_session_with_nobody_online_is_unchanged(self):
        s = self.session(men=10, women=12)
        self.assertEqual(s.total, 22)
        self.assertEqual(s.online_total, 0)

    def test_online_only_is_counted(self):
        """A service streamed with nobody in the building is still a
        service that happened."""
        s = self.session(online_men=5, online_women=6)
        self.assertEqual(s.in_person_total, 0)
        self.assertEqual(s.total, 11)

    def test_recording_saves_both(self):
        s = self.session()
        resp = self.client.post(f"/api/attendance-sessions/{s.id}/record/", {
            "men": 20, "women": 25, "online_men": 4, "online_women": 6,
            "new_comers": 3, "new_converts": 1,
        }, format="json")
        self.assertEqual(resp.status_code, 200)
        s.refresh_from_db()
        self.assertEqual(s.online_total, 10)
        self.assertEqual(s.total, 55)
        self.assertEqual(s.new_comers, 3)
        self.assertEqual(s.new_converts, 1)


class FellowshipTestCase(Base):
    """
    Several fellowships meet the same evening, so a session belongs to a
    fellowship as well as a date.
    """
    def setUp(self):
        super().setUp()
        self.women = Fellowship.objects.create(name="HCF Women", area="Riffa")
        self.men = Fellowship.objects.create(name="HCF Men", area="Manama")

    def test_two_fellowships_can_meet_on_the_same_date(self):
        a = self.session(self.house, fellowship=self.women, status="filled", women=9)
        b = self.session(self.house, fellowship=self.men, status="filled", men=8)
        self.assertEqual(a.date, b.date)
        self.assertNotEqual(a.fellowship, b.fellowship)

    def test_the_leader_is_recorded_per_session(self):
        """It changes week to week, so holding it on the fellowship would
        rewrite history every time somebody stood in."""
        sarah = Member.objects.create(surname="Osei", first_name="Sarah",
                                      location=self.bahrain, joined_date=self.today)
        grace = Member.objects.create(surname="Thomas", first_name="Grace",
                                      location=self.bahrain, joined_date=self.today)
        week1 = self.session(self.house, fellowship=self.women, led_by=sarah, status="filled")
        week2 = AttendanceSession.objects.create(
            meeting_type=self.house, date=self.today + datetime.timedelta(days=7),
            location=self.bahrain, mode="in-person", status="filled",
            fellowship=self.women, led_by=grace)
        self.assertNotEqual(week1.led_by, week2.led_by)

    def test_a_fellowship_with_sessions_cannot_be_deleted(self):
        """Those records would lose what they belong to."""
        self.session(self.house, fellowship=self.women, status="filled")
        resp = self.client.delete(f"/api/fellowships/{self.women.id}/")
        self.assertEqual(resp.status_code, 400)
        self.assertTrue(Fellowship.objects.filter(id=self.women.id).exists())

    def test_an_unused_fellowship_can_be_deleted(self):
        resp = self.client.delete(f"/api/fellowships/{self.men.id}/")
        self.assertEqual(resp.status_code, 204)

    def test_fellowships_are_listed_with_their_session_count(self):
        self.session(self.house, fellowship=self.women, status="filled")
        resp = self.client.get("/api/fellowships/")
        self.assertEqual(resp.status_code, 200)
        rows = {f["name"]: f["session_count"] for f in resp.data["results"]}
        self.assertEqual(rows["HCF Women"], 1)
        self.assertEqual(rows["HCF Men"], 0)


class SessionOfferingTestCase(Base):
    """
    A fellowship's offering is giving linked to the session, not amounts
    copied onto it. Recording it twice would mean two numbers that
    disagree the first time somebody corrects one.
    """
    def test_giving_can_belong_to_a_session(self):
        fellowship = Fellowship.objects.create(name="HCF Women")
        s = self.session(self.house, fellowship=fellowship, status="filled")
        fund = Fund.objects.create(name="Tithe")
        method = PaymentMethod.objects.create(name="Cash")
        Giving.objects.create(date=self.today, fund=fund, method=method,
                              amount=120, location=self.bahrain, session=s)
        self.assertEqual(s.giving.count(), 1)
        self.assertEqual(sum(g.amount for g in s.giving.all()), 120)

    def test_giving_no_longer_carries_a_destination(self):
        """Remittance happens once a month, not each time money is
        received, so asking per entry was the wrong question."""
        self.assertFalse(hasattr(Giving(), "remitted_to"))

    def test_giving_without_a_session_still_works(self):
        """Most giving is not tied to one meeting."""
        fund = Fund.objects.create(name="Tithe")
        method = PaymentMethod.objects.create(name="Cash")
        g = Giving.objects.create(date=self.today, fund=fund, method=method,
                                  amount=50, location=self.bahrain)
        self.assertIsNone(g.session)



class SpreadsheetTestCase(Base):
    def test_the_spreadsheet_downloads(self):
        self.session(status="filled", men=10, women=12)
        resp = self.client.get("/api/reports/spreadsheet/?year=%d&month=%d"
                               % (self.today.year, self.today.month))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertIn(".xlsx", resp["Content-Disposition"])

    def test_it_is_a_real_workbook_with_the_expected_sheets(self):
        import io
        from openpyxl import load_workbook
        self.session(status="filled", men=10, women=12)
        resp = self.client.get("/api/reports/spreadsheet/?year=%d&month=%d"
                               % (self.today.year, self.today.month))
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, ["Attendance", "Finance", "Fellowships"])

    def test_every_sheet_prints_as_one_page_across(self):
        """A sheet that spills its last columns onto a page of their own
        looks like a mistake to whoever opens it."""
        import io
        from openpyxl import load_workbook
        self.session(status="filled", men=10, women=12)
        resp = self.client.get("/api/reports/spreadsheet/?year=%d&month=%d"
                               % (self.today.year, self.today.month))
        wb = load_workbook(io.BytesIO(resp.content))
        for name in wb.sheetnames:
            self.assertTrue(wb[name].sheet_properties.pageSetUpPr.fitToPage, name)
